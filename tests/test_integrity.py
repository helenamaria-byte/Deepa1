"""The 17 required scenarios for decision integrity, quantities and rules.
Scenarios already covered elsewhere are referenced; the rest are tested here."""
import io
import re
from datetime import datetime

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.app.main import app, get_store
from src.app.status import plan_status
from src.engine.pipeline import run_engine
from src.jobs.run_plan import IST, run_plan
from src.rules_config import load_default_rules
from src.store.models import Override
from tests.conftest import RUN_DATE, StaticConnector, no_sleep
from tests.test_outputs_and_connectors import synthetic

RULES = load_default_rules()


def flat_sku(raw, sku, weekly, on_hand, perishable=None):
    """Give one SKU the same sales every week and a chosen stock figure."""
    s = raw["weekly_sales"]
    s.loc[s.SKU_ID == sku, "Units_Sold"] = str(weekly)
    raw["current_inventory"].loc[raw["current_inventory"].SKU_ID == sku, "Units_On_Hand"] = str(on_hand)
    if perishable is not None:
        raw["sku_master"].loc[raw["sku_master"].SKU_ID == sku, "Perishability"] = perishable


def no_festival(raw):
    cal = raw["calendar"]
    cal.loc[cal.Year_Month == "2026-10", "Is_Festival_Month"] = "FALSE"


def engine(raw, rules=RULES):
    return run_engine(raw, RUN_DATE, rules).recommendations.set_index("sku_id")


# 1-3: standard SKU below 6 / above 14 / within 6-14 weeks -> golden test (SKU003, SKU006, SKU005).
# 4: high demand > 20% -> golden test (SKU001). 6 and 8: perishable below 3 / above 4 -> golden (SKU004, SKU007).
# 11-12: insufficient history / stale or missing stock -> test_exceptions.py (NEW_LAUNCH, STALE_INVENTORY, ...).

@pytest.mark.parametrize("on_hand, expected", [(299, "STOCKOUT_RISK"), (300, "OK"), (700, "OK"),
                                               (701, "OVERSTOCK_RISK")])
def test_standard_window_boundaries(raw, on_hand, expected):
    no_festival(raw)
    flat_sku(raw, "SKU005", 50, on_hand)  # 50 a week: 300 = exactly 6 weeks, 700 = exactly 14 weeks
    assert engine(raw).loc["SKU005", "status"] == expected


def test_exactly_at_high_demand_threshold_is_not_flagged(raw):
    # 8 weeks at 75 then 4 at 100: avg12 = 83.33, 1.2 x avg12 = 100 = avg4 exactly.
    s = raw["weekly_sales"]
    for w in range(31, 43):
        s.loc[(s.SKU_ID == "SKU008") & (s.Week == f"2026-W{w}"), "Units_Sold"] = "75" if w < 39 else "100"
    res = engine(raw)
    assert res.loc["SKU008", "avg4"] == 100 and "HIGH_DEMAND" not in res.loc["SKU008", "flags"]
    s.loc[(s.SKU_ID == "SKU008") & (s.Week == "2026-W42"), "Units_Sold"] = "101"
    assert "HIGH_DEMAND" in engine(raw).loc["SKU008", "flags"]


@pytest.mark.parametrize("on_hand, expected", [(149, "STOCKOUT_RISK"), (150, "OK"), (175, "OK"), (200, "OK"),
                                               (201, "OVERSTOCK_RISK")])
def test_perishable_window_3_to_4_weeks(raw, on_hand, expected):
    no_festival(raw)
    flat_sku(raw, "SKU007", 50, on_hand)  # perishable; 150 = 3 weeks, 200 = 4 weeks
    res = engine(raw)
    assert (res.loc["SKU007", "min_cover"], res.loc["SKU007", "max_cover"]) == (3, 4)
    assert res.loc["SKU007", "status"] == expected


def test_festival_uplift_applied_once(raw):
    res = engine(raw)  # October: festival month
    r = res.loc["SKU002"]
    c = r["calc"]
    assert c["multiplier"] == 1.25 and c["baseline_demand"] == r["avg12"]
    assert c["weekly_demand"] == pytest.approx(r["avg12"] * 1.25, abs=0.01)
    assert r["cover_weeks"] == pytest.approx(r["on_hand"] / (r["avg12"] * 1.25), abs=0.01)  # not 1.25^2
    # The high-demand ratio uses raw averages: the uplift must not change it.
    rules = load_default_rules()
    rules.festival.enabled = False
    assert list(engine(raw, rules)["flags"].map(lambda f: "HIGH_DEMAND" in f)) == \
        list(res["flags"].map(lambda f: "HIGH_DEMAND" in f))


def pos(*rows):
    return pd.DataFrame([{"SKU_ID": s, "PO_Number": n, "Units_On_Order": str(u), "Expected_Date": d}
                         for s, n, u, d in rows])


def test_substantial_purchase_order_counted_only_if_in_time(raw):
    raw["open_purchase_orders"] = pos(("SKU001", "PO-A", 1000, "2026-10-29"),   # in time
                                      ("SKU003", "PO-B", 800, "2026-12-20"),    # after stock runs out
                                      ("SKU013", "PO-C", 60, ""))               # no date
    res = engine(raw)
    a = res.loc["SKU001"]
    assert a["on_order"] == 1000 and a["cover_weeks"] == pytest.approx(7.07, abs=0.01)
    assert a["status"] == "HIGH_DEMAND" and a["recommended_qty"] == 750  # rule 1 still applies
    assert a["calc"]["cover_after_order"] == pytest.approx((900 + 1000 + 750) / 268.65, abs=0.01)
    assert "inside the 6 to 14-week range" in a["calc"]["checks"][0]
    b = res.loc["SKU003"]
    assert b["on_order"] == 0 and b["calc"]["on_order_excluded"] == 800 and b["status"] == "STOCKOUT_RISK"
    assert "after stock on hand runs out" in b["calc"]["incoming"][0]["note"]
    assert "still below the 6-week minimum" in b["calc"]["checks"][0]
    assert res.loc["SKU013", "calc"]["incoming"][0]["counted"] is False


def test_counting_all_orders_is_a_rule_switch(raw):
    raw["open_purchase_orders"] = pos(("SKU003", "PO-B", 800, "2026-12-20"))
    rules = load_default_rules()
    rules.open_orders.count_only_if_arrives_before_stockout = False
    assert engine(raw, rules).loc["SKU003", "on_order"] == 800


# ---- decisions ---------------------------------------------------------------
@pytest.fixture
def plan(raw, monkeypatch):
    store = get_store()
    run = run_plan(RUN_DATE, store=store, connector=StaticConnector(raw), sleep=no_sleep, send_email=False)
    client = TestClient(app)
    client.post("/me", data={"who": "Raghav"})
    return store, run, client


def test_repeated_approval_is_recorded_once(plan):
    store, run, c = plan
    for _ in range(3):
        c.post(f"/sku/{run['id']}/SKU003", data={"decision": "APPROVED", "next": "/"})
    events = store.overrides(run["id"], include_deleted=True)
    assert len(events) == 1 and store.decision_progress(run["id"])["decided"] == 1


def test_change_undo_and_history(plan):
    store, run, c = plan
    rid = run["id"]
    c.post(f"/sku/{rid}/SKU003", data={"decision": "APPROVED", "next": "/"})
    c.post(f"/sku/{rid}/SKU003", data={"decision": "CHANGED", "final_qty": "350", "next": "/"})  # small: no reason
    cur = store.latest_decisions(rid)["SKU003"]
    assert cur["final_qty"] == 350 and cur["previous_qty"] == 300 and cur["state"] == "current"
    assert cur["reason"] == "Quantity changed (no reason given)"
    # A material change needs a reason.
    r = c.post(f"/sku/{rid}/SKU003", data={"decision": "CHANGED", "final_qty": "600", "next": "/"})
    assert "Please add a reason" in r.text and store.latest_decisions(rid)["SKU003"]["final_qty"] == 350
    c.post(f"/sku/{rid}/SKU003", data={"decision": "CHANGED", "final_qty": "600", "reason": "Festival promo",
                                       "next": "/"})
    cur = store.latest_decisions(rid)["SKU003"]
    assert cur["final_qty"] == 600 and cur["previous_qty"] == 350
    # Undo the current one: the SKU is pending again and nothing is erased.
    r = c.post(f"/decisions/{cur['id']}/delete", data={"next": "/"})
    assert "Undone" in r.text and "SKU003" not in store.latest_decisions(rid)
    states = [o["state"] for o in store.overrides(rid, include_deleted=True)]
    assert states == ["undone", "replaced", "replaced"]
    # A replaced decision cannot be undone.
    old = store.overrides(rid, include_deleted=True)[-1]
    assert not store.delete_override(old["id"], "x")
    history = c.get("/decisions").text
    assert "Replaced" in history and "Undone" in history and "350 → <b>600</b>" in history


def test_counts_agree_across_pages(plan):
    store, run, c = plan
    rid = run["id"]
    for sku in ("SKU001", "SKU002", "SKU006"):
        c.post(f"/sku/{rid}/{sku}", data={"decision": "APPROVED", "next": "/"})
    c.post(f"/sku/{rid}/SKU001", data={"decision": "APPROVED", "next": "/"})  # repeat click
    assert store.decision_progress(rid) == {"total": 15, "decided": 3, "pending": 12}
    assert "3 of 15 done" in c.get("/").text
    assert "Current plan #%d (Thu 22 Oct): 3 of 15 decided" % rid in c.get("/decisions").text
    assert ">3 of 15<" in c.get("/analytics").text
    assert "Approve all 12 SKUs still waiting?" in c.get("/recommendations?show=todo").text or \
        "Approve all" in c.get("/recommendations?show=todo").text
    todo = c.get("/recommendations?show=todo").text
    assert len(re.findall(r'<tr id="SKU\d+">', todo)) == 12


def test_old_duplicate_clicks_are_marked_replaced_not_deleted(plan):
    store, run, _ = plan
    with store.Session.begin() as s:  # the live case: three identical approvals stored before v4
        for _ in range(3):
            s.add(Override(run_id=run["id"], sku_id="SKU003", decision="APPROVED", recommended_qty=300,
                           final_qty=300, reason="Good to go", username="Raghav"))
    store.migrate()
    events = store.overrides(run["id"], include_deleted=True)
    assert len(events) == 3 and [e["state"] for e in events].count("current") == 1
    assert store.decision_progress(run["id"])["decided"] == 1


# ---- Excel import --------------------------------------------------------------
def sheet(c, run_id, edits, run_label=None):
    wb = load_workbook(io.BytesIO(c.get(f"/runs/{run_id}/workbook").content))
    ws = wb["Override log"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    used = set()
    for sku, decision, qty, reason in edits:
        if sku in rows and sku not in used:
            row = rows[sku]
            used.add(sku)
        else:  # a second row for the same SKU (duplicate) or a SKU not in the sheet
            ws.append([sku])
            row = list(ws.iter_rows(min_row=ws.max_row))[0]
        row[3].value, row[4].value, row[5].value = decision, qty, reason
    if run_label is not None:
        for r in wb["Run info"].iter_rows(min_row=2):
            if r[0].value == "Run ID":
                r[1].value = run_label
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_import_all_or_nothing(plan):
    store, run, c = plan
    rid = run["id"]
    bad = sheet(c, rid, [("SKU013", "Approve", None, None), ("SKU013", "Approve", None, None),
                         ("SKU003", "Change", "lots", "x")])
    r = c.post("/decisions/upload", files={"file": ("p.xlsx", bad)})
    assert "Nothing was imported" in r.text and "more than once" in r.text and "whole number" in r.text
    assert store.decision_progress(rid)["decided"] == 0
    wrong_plan = sheet(c, rid, [("SKU013", "Approve", None, None)], run_label=999)
    assert "is for plan #999" in c.post("/decisions/upload", files={"file": ("p.xlsx", wrong_plan)}).text
    good = sheet(c, rid, [("SKU013", "Approve", None, None), ("SKU003", "Change", 320, "")])
    r = c.post("/decisions/upload", files={"file": ("p.xlsx", good)})
    assert "Imported 2 decisions." in r.text
    r = c.post("/decisions/upload", files={"file": ("p.xlsx", good)})  # same file again
    assert "Imported 0 decisions." in r.text and "2 rows matched decisions already recorded" in r.text
    assert len(store.overrides(rid, include_deleted=True)) == 2


# ---- plan status and full dataset ---------------------------------------------------
def test_plan_status_reflects_real_inputs(plan):
    store, run, c = plan
    home = c.get("/").text
    assert "Sales data covers" in home and "2026-W27 to 2026-W42" in home
    assert "15 processed (15 in the SKU master)" in home and "Safe to act on" in home
    from src.connectors.factory import make_connector
    from src.jobs.run_plan import source_label
    assert "demo data, not a live SAP export" in source_label(make_connector(""))  # bundled sample files
    recs = store.recommendations(run["id"])
    later = datetime(2026, 11, 6, 9, 0, tzinfo=IST)  # stock figures 18 days old by then
    status = plan_status(store.get_run(run["id"]), recs, RULES, later, later)
    assert status["level"] == "unsafe" and any("days old" in w for w in status["warnings"])


def test_full_dataset_is_processed_without_silent_omissions():
    data = synthetic(n_skus=400)
    data["current_inventory"] = data["current_inventory"].iloc[5:]           # 5 SKUs with no stock row
    extra = data["weekly_sales"][data["weekly_sales"].SKU_ID == "S0000"].assign(SKU_ID="ORPHAN1")
    data["weekly_sales"] = pd.concat([data["weekly_sales"], extra])          # 1 SKU with no master record
    store = get_store()
    run = run_plan(RUN_DATE, store=store, connector=StaticConnector(data), sleep=no_sleep, send_email=False)
    recs = store.recommendations(run["id"])
    assert len(recs) == 401 == run["data_info"]["skus_processed"] and run["data_info"]["skus_in_master"] == 400
    missing = [r for r in recs if "MISSING_INVENTORY" in r["exception_codes"]]
    orphan = [r for r in recs if "ORPHAN_SKU" in r["exception_codes"]]
    assert len(missing) == 5 and len(orphan) == 1
    assert all(r["recommended_qty"] is None for r in missing + orphan)
