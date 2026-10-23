"""Workbook, emails, connector swap, Zoho stub and the scale test."""
import io
import time
from datetime import date

import httpx
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.connectors.csv_folder import CsvFolderConnector
from src.connectors.factory import make_connector
from src.connectors.https import HttpsConnector
from src.connectors.zoho import OPEN_PO_COLUMNS, ZohoConnector
from src.engine.pipeline import run_engine
from src.jobs.run_plan import run_plan
from src.jobs.weekly_summary import send_weekly_summary
from src.outputs.email import Mailer
from src.outputs.workbook import SHEETS, build_workbook, read_override_sheet
from src.rules_config import load_default_rules
from src.settings import ROOT
from src.store.repo import Store
from tests.conftest import RUN_DATE, StaticConnector, no_sleep


def make_run(raw):
    store = Store()
    run = run_plan(RUN_DATE, store=store, connector=StaticConnector(raw), sleep=no_sleep)
    return store, run


def test_workbook_has_five_sheets_and_opens(raw):
    store, run = make_run(raw)
    data = build_workbook(run, store.recommendations(run["id"]), [], load_default_rules(), 1)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == SHEETS
    rec_rows = list(wb["Recommendations"].iter_rows(values_only=True))
    assert len(rec_rows) == 16 and rec_rows[0][0] == "SKU_ID"


def test_override_sheet_round_trip(raw):
    store, run = make_run(raw)
    data = build_workbook(run, store.recommendations(run["id"]), [], load_default_rules(), 1)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Override log"]
    row = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "SKU003")
    row[3].value, row[4].value, row[5].value = "Change", 450, "Distributor promo"
    buf = io.BytesIO()
    wb.save(buf)
    parsed = read_override_sheet(buf.getvalue())
    assert parsed == [{"sku_id": "SKU003", "decision": "CHANGED", "final_qty": 450, "reason": "Distributor promo"}]


def test_three_emails_rendered_to_outbox(raw, isolated_env):
    store, run = make_run(raw)  # alert
    assert send_weekly_summary(store, end=date(2026, 10, 23)) == "OUTBOX"
    from src.jobs.run_plan import send_failure
    send_failure(Mailer(store), 99, RUN_DATE, 3, "Traceback...\nConnectionError: boom")
    outbox = isolated_env / "var" / "outbox"
    for kind in ("alert", "weekly_summary", "failure"):
        assert list(outbox.glob(f"*-{kind}.eml")) and list(outbox.glob(f"*-{kind}.html")), kind
    alert = next(outbox.glob("*-alert.html")).read_text(encoding="utf-8")
    assert "7 SKUs need review" in alert and "SKU001" in alert
    eml = next(outbox.glob("*-alert.eml")).read_bytes()
    assert b".xlsx" in eml  # workbook attached
    kinds = {e["kind"]: e for e in store.emails()}
    assert kinds["alert"]["recipients"] == "raghav@example.com, procurement@example.com"
    assert kinds["weekly_summary"]["recipients"] == "deepa@example.com"


def test_resend_payload(raw, monkeypatch):
    monkeypatch.setenv("RESEND_API_KEY", "re_test")
    seen = {}

    def handler(request):
        seen["json"] = request.read()
        seen["auth"] = request.headers["authorization"]
        return httpx.Response(200, json={"id": "abc"})

    store = Store()
    store.migrate()
    status = Mailer(store, transport=httpx.MockTransport(handler)).send(
        "alert", ["a@example.com"], "Hi", "<p>x</p>", attachments=[("plan.xlsx", b"123")])
    assert status == "SENT" and seen["auth"] == "Bearer re_test"
    assert b'"filename":"plan.xlsx"' in seen["json"].replace(b" ", b"")


def test_https_connector_swaps_in_with_identical_results(raw):
    def handler(request):
        name = request.url.path.rsplit("/", 1)[-1]
        return httpx.Response(200, content=(ROOT / "data" / name).read_bytes())

    https = HttpsConnector("https://sap.example.com/exports", token="t",
                           transport=httpx.MockTransport(handler)).fetch()
    a = run_engine(https, RUN_DATE, load_default_rules()).recommendations
    b = run_engine(CsvFolderConnector(ROOT / "data").fetch(), RUN_DATE, load_default_rules()).recommendations
    pd.testing.assert_frame_equal(a, b)


def test_factory_selects_connector_from_config():
    assert make_connector("").name == "csv_folder"
    assert make_connector("https://x.example.com/exp").name == "https"
    assert make_connector("s3://bucket/prefix").name == "s3"
    assert make_connector("sftp://u@host/dir").name == "sftp"


def test_zoho_stub_behind_flag():
    z = ZohoConnector(enabled=False)
    assert z.enabled is False
    pos = z.fetch_open_pos()
    assert list(pos.columns) == OPEN_PO_COLUMNS and pos.empty


def synthetic(n_skus=400, n_weeks=156, n_dist=1200, seed=7):
    rng = np.random.default_rng(seed)
    weeks = pd.date_range("2023-10-23", periods=n_weeks, freq="W-MON")
    labels = [f"{d.isocalendar().year}-W{d.isocalendar().week:02d}" for d in weeks]
    skus = [f"S{i:04d}" for i in range(n_skus)]
    base = rng.integers(20, 400, n_skus)
    sales = pd.DataFrame({
        "SKU_ID": np.repeat(skus, n_weeks), "Week": np.tile(labels, n_skus),
        "Units_Sold": (np.repeat(base, n_weeks) * rng.uniform(0.85, 1.15, n_skus * n_weeks)).astype(int).astype(str)})
    dists = [f"D{i:04d}" for i in range(n_dist)]
    master = pd.DataFrame({"SKU_ID": skus, "SKU_Name": skus, "Category": "Cat",
                           "Perishability": rng.choice(["High", "Low"], n_skus),
                           "Standard_Reorder_Qty": (base * 3).astype(str), "Unit_of_Measure": "Units",
                           "Distributor_ID": rng.choice(dists, n_skus)})
    inv = pd.DataFrame({"SKU_ID": skus, "Location": "WH", "Units_On_Hand": (base * rng.uniform(1, 16, n_skus)).astype(int).astype(str),
                        "As_Of_Date": "2026-10-19"})
    dist = pd.DataFrame({"Distributor_ID": dists, "Distributor_Name": dists, "Region": "R", "Status": "Active"})
    cal = pd.read_csv(ROOT / "data" / "calendar.csv", dtype=str)
    return {"weekly_sales": sales, "current_inventory": inv, "sku_master": master,
            "distributor_master": dist, "calendar": cal}


def test_scale_400_skus_156_weeks_under_60_seconds():
    data = synthetic()
    assert len(data["weekly_sales"]) == 400 * 156
    start = time.perf_counter()
    run = run_plan(RUN_DATE, store=Store(), connector=StaticConnector(data), sleep=no_sleep, send_email=False)
    elapsed = time.perf_counter() - start
    assert run["status"] == "SUCCESS" and run["counts"]["TOTAL"] == 400
    assert elapsed < 60, elapsed
    print(f"scale run: {elapsed:.1f}s")
