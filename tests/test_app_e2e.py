"""End-to-end over HTTP: login, every page, override round trip, Excel upload,
settings versioning, cron protection."""
import io
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.app.main import app, get_store
from src.jobs.bootstrap import bootstrap
from src.jobs.run_plan import run_plan
from tests.conftest import RUN_DATE, StaticConnector, no_sleep


@pytest.fixture
def client(raw, monkeypatch):
    monkeypatch.setenv("DASHBOARD_USERS", "raghav@example.com:pw-raghav")
    store = get_store()
    bootstrap(store)
    run_plan(RUN_DATE, store=store, connector=StaticConnector(raw), sleep=no_sleep)
    c = TestClient(app)
    r = c.post("/login", data={"username": "raghav@example.com", "password": "pw-raghav"}, follow_redirects=False)
    assert r.status_code == 303
    return c


def test_pages_require_login_when_auth_on(monkeypatch):
    monkeypatch.setenv("AUTH_REQUIRED", "true")
    c = TestClient(app)
    for path in ["/", "/decisions", "/analytics", "/runs", "/settings"]:
        r = c.get(path, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"].startswith("/login"), path
    assert c.post("/runs/new", follow_redirects=False).status_code == 303


@pytest.fixture
def open_client(raw):
    store = get_store()
    run_plan(RUN_DATE, store=store, connector=StaticConnector(raw), sleep=no_sleep)
    return TestClient(app)


def test_pages_open_without_login(open_client):
    for path in ["/", "/decisions", "/analytics", "/runs", "/settings", "/sku/1/SKU001"]:
        r = open_client.get(path, follow_redirects=False)
        assert r.status_code == 200, path
    assert "Who is reviewing?" in open_client.get("/").text
    assert 'name="who"' in open_client.get("/sku/1/SKU001").text


def test_open_mode_decision_needs_name_and_remembers_it(open_client):
    r = open_client.post("/sku/1/SKU003", data={"decision": "APPROVED", "reason": "Fine"})
    assert "enter your name" in r.text
    assert get_store().overrides(1) == []
    r = open_client.post("/sku/1/SKU003", data={"decision": "APPROVED", "reason": "Fine", "who": "  Raghav  K "})
    assert "Saved your decision for SKU003" in r.text
    assert get_store().overrides(1)[0]["username"] == "Raghav K"
    # The name is remembered and pre-filled, and used when the form omits it.
    assert 'value="Raghav K"' in open_client.get("/sku/1/SKU001").text
    r = open_client.post("/sku/1/SKU001", data={"decision": "APPROVED", "reason": "Agree"})
    assert "Saved your decision for SKU001" in r.text
    assert get_store().latest_decisions(1)["SKU001"]["username"] == "Raghav K"


def test_wrong_password_rejected(client):
    c = TestClient(app)
    r = c.post("/login", data={"username": "raghav@example.com", "password": "nope"})
    assert "password don" in r.text and "t match" in r.text
    assert "fr_session" not in r.cookies


def test_every_page_loads(client):
    for path in ["/", "/?show=todo", "/?status=STOCKOUT_RISK",
                 "/decisions", "/analytics", "/runs", "/settings", "/sku/1/SKU001", "/judgment-log.md", "/health"]:
        r = client.get(path)
        assert r.status_code == 200, path
    home = client.get("/").text
    assert "0 of 15 done" in home and "Reorder now" in home
    rec = client.get("/?status=STOCKOUT_RISK").text
    assert "SKU001" in rec and "750" in rec and "SKU006" not in rec


def test_override_round_trip_and_delete(client):
    r = client.post("/sku/1/SKU003", data={"decision": "CHANGED", "final_qty": "420",
                                            "reason": "HTTP test override"})
    assert "Saved your decision for SKU003" in r.text
    decisions = client.get("/decisions").text
    assert "HTTP test override" in decisions and "raghav@example.com" in decisions
    oid = re.search(r"/decisions/(\d+)/delete", decisions).group(1)
    r = client.post(f"/decisions/{oid}/delete")
    assert "Undone" in r.text
    assert get_store().latest_decisions(1) == {}
    assert get_store().overrides(1, include_deleted=True)[0]["state"] == "undone"  # kept in history


def test_approve_all(client):
    r = client.post("/approve-all", data={"run_id": "1", "reason": "Reviewed with procurement"})
    assert "Approved 15 suggestions." in r.text
    assert "All 15 decisions made" in client.get("/").text


def test_excel_download_and_upload(client):
    r = client.get("/runs/1/workbook")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Override log"]
    row = next(x for x in ws.iter_rows(min_row=2) if x[0].value == "SKU013")
    row[3].value, row[5].value = "Approve", "Looks right"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/decisions/upload", files={"file": ("plan.xlsx", buf.getvalue())})
    assert "Imported 1 decisions." in r.text
    assert get_store().latest_decisions(1)["SKU013"]["final_qty"] == 120


def test_settings_save_creates_version(client):
    page = client.get("/settings").text
    form = dict(re.findall(r'name="(\w+)" value="([^"]+)"', page))
    form["festival_enabled"] = "on"
    form["gen_min"] = "5"
    form["note"] = "test"
    r = client.post("/settings", data=form)
    assert "Saved as version 2" in r.text
    assert get_store().current_rules()[1].cover.general.min_weeks == 5


def test_settings_rejects_bad_values(client):
    page = client.get("/settings").text
    form = dict(re.findall(r'name="(\w+)" value="([^"]+)"', page))
    form["gen_min"], form["gen_max"] = "20", "10"
    assert "Not saved" in client.post("/settings", data=form).text


def test_run_now_button(client):
    r = client.post("/runs/new")
    assert r.status_code == 200
    assert len(get_store().list_runs()) == 2


def test_cron_requires_secret(client):
    assert TestClient(app).get("/api/cron/run").status_code == 401
    assert TestClient(app).get("/api/cron/run", headers={"Authorization": "Bearer wrong"}).status_code == 401
    r = TestClient(app).get("/api/cron/weekly", headers={"Authorization": "Bearer test-cron"})
    assert r.status_code == 200 and r.json()["weekly_summary"] == "OUTBOX"


def test_landing_page_is_raghavs_action_list(client):
    home = client.get("/").text
    # Sections in the order he acts on them; OK SKUs are folded away.
    assert home.index("Reorder now") < home.index("Pause the next order") < home.index("No order needed")
    # OK SKUs are a section of their own, each with Approve and Change in place.
    assert home.count("Approve: no order") == 8 and home.count("Change quantity") == 15
    assert "Approve 750" in home and "Approve pause" in home
    # The rule behind each flag, in plain words with the SKU's own numbers.
    assert "Sales +21% vs 12-week average (trigger: over +20%)" in home
    assert "3.35 weeks of stock, below the 6-week minimum" in home
    assert "4.08 weeks of stock, above the 4-week perishable maximum" in home
    assert "How the flags are decided" in home and "Festival months" in home


def test_one_click_approve_and_undo(client):
    r = client.post("/sku/1/SKU001", data={"decision": "APPROVED", "next": "/#SKU001"})
    assert "Saved your decision for SKU001" in r.text and "1 of 15 done" in r.text
    d = get_store().latest_decisions(1)["SKU001"]
    assert d["final_qty"] == 750 and d["reason"] == "Approved as recommended"
    oid = re.search(r'/decisions/(\d+)/delete" class="inline">\s*<input type="hidden" name="next" value="/#SKU001"',
                    r.text).group(1)
    r = client.post(f"/decisions/{oid}/delete", data={"next": "/#SKU001"})
    assert "0 of 15 done" in r.text


def test_change_reason_is_optional(client):
    r = client.post("/sku/1/SKU003", data={"decision": "CHANGED", "final_qty": "350", "next": "/"})
    assert "Saved your decision for SKU003" in r.text and "Changed to 350 units" in r.text
    d = get_store().latest_decisions(1)["SKU003"]
    assert d["final_qty"] == 350 and d["reason"] == "Quantity changed (no reason given)"
    # The SKU007 case from review: a pause changed to 3 units (standard 400) needs no reason.
    r = client.post("/sku/1/SKU007", data={"decision": "CHANGED", "final_qty": "3", "next": "/"})
    assert "Saved your decision for SKU007" in r.text
    r = client.post("/sku/1/SKU013", data={"decision": "CHANGED", "final_qty": "90", "reason": "Promo", "next": "/"})
    assert get_store().latest_decisions(1)["SKU013"]["reason"] == "Promo"


def test_menu_has_no_exceptions_or_runs_tab(client):
    home = client.get("/").text
    nav = home[home.index("<nav>"):home.index("</nav>")]
    assert "Exceptions" not in nav and "Runs" not in nav
    for tab in ("Today", "Decisions", "Analytics", "Settings"):
        assert tab in nav
    r = client.get("/exceptions", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/"


def test_approve_all_from_landing_page_needs_no_reason(client):
    r = client.post("/recommendations/approve-all", data={"run_id": "1", "next": "/"})
    assert "Approved 15 suggestions." in r.text and "All 15 decisions made" in r.text


def test_ok_sku_can_be_approved_or_changed_in_place(client):
    r = client.post("/sku/1/SKU002", data={"decision": "APPROVED", "next": "/#SKU002"})
    assert "✓ Approved · no order" in r.text and "1 of 15 done" in r.text
    d = get_store().latest_decisions(1)["SKU002"]
    assert d["decision"] == "APPROVED" and d["final_qty"] is None and d["recommended_qty"] is None
    r = client.post("/sku/1/SKU005", data={"decision": "CHANGED", "final_qty": "100", "reason": "Diwali promo",
                                          "next": "/?status=OK#SKU005"})
    assert "Saved your decision for SKU005" in r.text and "Changed to 100 units" in r.text
    rec = client.get("/?status=OK").text
    assert "Approve: no order" in rec and "No order" in rec
